"""ATS Resume Checker - Streamlit + Google Gemini Flash.

Upload a resume (PDF or DOCX), optionally paste a job description, and get:
  * an ATS score (0-100) with a category breakdown
  * strengths, weaknesses and concrete improvements
  * missing keywords and rewritten bullet-point suggestions
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = "gemini-2.5-flash"
MAX_RESUME_CHARS = 20000  # keeps prompts small and fast
MIN_RESUME_CHARS = 200  # below this the file is probably scanned / empty

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and
professional resume coach. You evaluate resumes strictly and honestly.
Never invent experience that is not in the resume.
Respond ONLY with valid JSON that follows the schema you are given."""

JSON_SCHEMA_HINT = """{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "breakdown": {
    "formatting_and_structure": <integer 0-100>,
    "keywords_and_skills": <integer 0-100>,
    "work_experience_impact": <integer 0-100>,
    "education_and_certifications": <integer 0-100>,
    "readability_and_grammar": <integer 0-100>
  },
  "job_match_score": <integer 0-100 or null if no job description was given>,
  "strengths": ["..."],
  "weaknesses": ["..."],
  "improvements": [
    {"priority": "High|Medium|Low", "issue": "...", "fix": "..."}
  ],
  "missing_keywords": ["..."],
  "rewrite_suggestions": [
    {"original": "...", "improved": "..."}
  ]
}"""


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # resumes often keep content in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# --------------------------------------------------------------------------
# Quick rule-based checks (no AI needed)
# --------------------------------------------------------------------------
def basic_checks(text: str) -> list[tuple[bool, str]]:
    lower = text.lower()
    return [
        (bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)), "Email address found"),
        (bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)), "Phone number found"),
        (
            bool(re.search(r"linkedin\.com|github\.com", lower)),
            "LinkedIn / GitHub link found",
        ),
        (
            any(k in lower for k in ("experience", "employment", "work history")),
            "Experience section found",
        ),
        (any(k in lower for k in ("education", "university", "college")), "Education section found"),
        (any(k in lower for k in ("skills", "technologies", "tools")), "Skills section found"),
        (bool(re.search(r"\d+\s?%|\$\s?\d|\b\d{2,}\b", text)), "Quantified achievements (numbers) found"),
        (300 <= len(text.split()) <= 1200, "Length is in a healthy range (300-1200 words)"),
    ]


# --------------------------------------------------------------------------
# Gemini call
# --------------------------------------------------------------------------
def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences / extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI returned an unreadable response. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure every field exists with the right type so the UI never crashes."""
    breakdown = data.get("breakdown") or {}
    jm = data.get("job_match_score")
    return {
        "ats_score": _clamp(data.get("ats_score")),
        "summary": str(data.get("summary") or ""),
        "breakdown": {k: _clamp(v) for k, v in breakdown.items()},
        "job_match_score": None if jm is None else _clamp(jm),
        "strengths": [str(x) for x in data.get("strengths") or []],
        "weaknesses": [str(x) for x in data.get("weaknesses") or []],
        "improvements": [
            i for i in (data.get("improvements") or []) if isinstance(i, dict)
        ],
        "missing_keywords": [str(x) for x in data.get("missing_keywords") or []],
        "rewrite_suggestions": [
            r for r in (data.get("rewrite_suggestions") or []) if isinstance(r, dict)
        ],
    }


def analyze_resume(client, resume_text: str, job_description: str = "") -> dict:
    jd_part = (
        f"\n\nJOB DESCRIPTION:\n{job_description.strip()[:8000]}\n\n"
        "Also score how well the resume matches this job (job_match_score) and list "
        "keywords from the job description that are missing from the resume."
        if job_description.strip()
        else "\n\nNo job description was provided. Set job_match_score to null and "
        "judge keywords against common standards for the candidate's apparent field."
    )
    prompt = (
        "Analyze the resume below for ATS compatibility and quality.\n"
        "Give 4-8 improvements ordered by priority and 2-4 rewrite suggestions that "
        "improve real bullet points from the resume (action verbs, metrics).\n"
        f"Return JSON exactly in this shape:\n{JSON_SCHEMA_HINT}"
        f"{jd_part}\n\nRESUME:\n{resume_text[:MAX_RESUME_CHARS]}"
    )
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


# --------------------------------------------------------------------------
# UI helpers
# --------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def get_api_key() -> str:
    # 1) Streamlit secrets (used on Streamlit Community Cloud)
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally - that's fine
    # 2) Environment variable
    return os.getenv("GEMINI_API_KEY", "")


def render_results(result: dict, resume_text: str) -> None:
    score = result["ats_score"]
    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{score}/100", score_label(score))
        st.progress(score / 100)
        if result["job_match_score"] is not None:
            st.metric("Job Match", f"{result['job_match_score']}/100")
    with col2:
        st.subheader("Summary")
        st.write(result["summary"])

    if result["breakdown"]:
        st.subheader("Score breakdown")
        for key, val in result["breakdown"].items():
            st.write(key.replace("_", " ").title())
            st.progress(val / 100, text=f"{val}/100")

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result["strengths"] or ["None listed"]:
            st.success(s)
    with right:
        st.subheader("Weaknesses")
        for w in result["weaknesses"] or ["None listed"]:
            st.warning(w)

    st.subheader("How to improve")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for item in result["improvements"]:
        pr = str(item.get("priority", "Medium"))
        with st.expander(f"{icons.get(pr.lower(), '🟠')} {pr}: {item.get('issue', '')}"):
            st.write(item.get("fix", ""))

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(" ".join(f"`{k}`" for k in result["missing_keywords"]))

    if result["rewrite_suggestions"]:
        st.subheader("Suggested rewrites")
        for r in result["rewrite_suggestions"]:
            st.markdown(f"**Before:** {r.get('original', '')}")
            st.markdown(f"**After:** {r.get('improved', '')}")
            st.write("")

    report = json.dumps(result, indent=2)
    st.download_button(
        "Download report (JSON)", report, "ats_report.json", "application/json"
    )
    with st.expander("View extracted resume text"):
        st.text(resume_text)


# --------------------------------------------------------------------------
# Main app
# --------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if api_key:
            st.success("Gemini API key loaded")
        else:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        st.markdown(f"Model: `{MODEL_NAME}`")
        st.info("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
    job_description = st.text_area(
        "Job description (optional, for a targeted score)",
        height=150,
        placeholder="Paste the job description here to see how well you match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_resume_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text was found. If your resume is a scanned image, "
                "ATS systems can't read it either - export a text-based PDF or DOCX."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                client = genai.Client(api_key=api_key)
                result = analyze_resume(client, text, job_description)
        except Exception as exc:
            msg = str(exc)
            if "API key" in msg or "API_KEY" in msg or "403" in msg or "400" in msg:
                st.error("Gemini rejected the request. Check that your API key is valid.")
            elif "429" in msg or "quota" in msg.lower():
                st.error("Rate limit reached. Wait a minute and try again.")
            else:
                st.error(f"Analysis failed: {msg}")
            return

        # Keep results in session_state so they survive Streamlit reruns
        # (e.g. clicking the download button).
        st.session_state["report"] = {
            "result": result,
            "text": text,
            "checks": basic_checks(text),
        }

    report = st.session_state.get("report")
    if report:
        st.subheader("Quick checks")
        cols = st.columns(2)
        for i, (ok, label) in enumerate(report["checks"]):
            cols[i % 2].write(("✅ " if ok else "❌ ") + label)
        render_results(report["result"], report["text"])


if __name__ == "__main__":
    main()
